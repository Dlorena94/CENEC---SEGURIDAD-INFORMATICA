import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl import Workbook
from io import BytesIO

eventos = [
    ("09:45", "10:40", "Smart Shortcuts: Google AI for a More Productive You", "Charla", "Sala A", "Ponente 1"),
    ("09:45", "11:40", "Reconocimiento de imágenes con Google Cloud y Raspberry", "Taller", "Sala B", "Ponente 2"),
    ("09:45", "11:40", "MCP: Securing the unsecure", "Taller", "Sala C", "Ponente 3"),
    ("10:45", "11:40", "¿El principio del fin? Peligros y oportunidades", "Charla", "Sala A", "Ponente 4"),
    ("10:45", "11:40", "Beyond the Code: Thriving as a Senior Engineer in the Age of Ownership", "Charla", "Sala B", "Ponente 5"),
    ("12:30", "13:25", "La observabilidad es mi pastor, nada me falta", "Charla", "Sala A", "Ponente 6"),
    ("12:30", "13:25", "From Android to Full‑Stack: A Developer's Journey to the Great Unknown", "Charla", "Sala B", "Ponente 7"),
    ("12:30", "14:25", "Building your AI agent with Agent Development Kit (ADK)", "Taller", "Sala C", "Ponente 8"),
    ("12:30", "14:25", "Manual para romper tu sistema: Antipatrones de diseño", "Taller", "Sala D", "Ponente 9"),
    ("13:30", "14:25", "Entendiendo el Presente con los Fundamentos del Software", "Charla", "Sala A", "Ponente 10"),
    ("13:30", "14:25", "Tenemos que hablar... (de salud mental en el mundo tech)", "Charla", "Sala B", "Ponente 11"),
    ("16:00", "16:55", "How to build a diagnostic agent with Vertex AI SDK and osquery", "Charla", "Sala A", "Ponente 12"),
    ("16:00", "16:55", "Kotlin MultiPlatform: La thermomix del desarrollo", "Charla", "Sala B", "Ponente 13"),
    ("16:00", "17:55", "Creando un prototipo de producto partiendo de cero", "Taller", "Sala C", "Ponente 14"),
    ("16:00", "17:55", "Coding for emotion. For diverse minds and emotional realities", "Taller", "Sala D", "Ponente 15"),
    ("17:00", "17:55", "IA con salero: descubre MCP Servers en Málaga", "Charla", "Sala A", "Ponente 16"),
    ("18:10", "19:00", "CQRS Práctico", "Charla", "Sala A", "Ponente 17"),
    ("09:00", "09:30", "Inauguración", "Evento", "Auditorio", "Organización"),
    ("11:40", "12:30", "Coffee Break", "Descanso", "Zona Café", ""),
    ("14:25", "16:00", "Comida", "Descanso", "Zona Comedor", ""),
    ("17:55", "18:10", "Coffee Break", "Descanso", "Zona Café", ""),
    ("19:00", "19:30", "Clausura", "Evento", "Auditorio", "Organización"),
]

df = pd.DataFrame(eventos, columns=["Hora inicio", "Hora fin", "Título", "Tipo", "Sala", "Ponente"])

def time_to_minutes(t):
    h, m = map(int, t.split(':'))
    return h * 60 + m

def detect_solapamientos(df):
    n = len(df)
    solapados = [False] * n
    for i in range(n):
        start_i = time_to_minutes(df.loc[i, "Hora inicio"])
        end_i = time_to_minutes(df.loc[i, "Hora fin"])
        for j in range(i + 1, n):
            start_j = time_to_minutes(df.loc[j, "Hora inicio"])
            end_j = time_to_minutes(df.loc[j, "Hora fin"])
            if not (end_i <= start_j or end_j <= start_i):
                solapados[i] = True
                solapados[j] = True
    return solapados

df["Solapado"] = detect_solapamientos(df)

wb = Workbook()
ws1 = wb.active
ws1.title = "Programa Completo"
headers = ["Hora", "Título", "Tipo", "Sala", "Ponente"]
ws1.append(headers)

from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

header_fill = PatternFill("solid", fgColor="003366")
header_font = Font(color="FFFFFF", bold=True)
center_align = Alignment(horizontal="center", vertical="center")
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
red_fill = PatternFill("solid", fgColor="FFC7CE")

for col_num, _ in enumerate(headers, 1):
    cell = ws1.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = border

for idx, row in df.iterrows():
    hora = f"{row['Hora inicio']} - {row['Hora fin']}"
    titulo = row['Título']
    tipo = row['Tipo']
    sala = row['Sala']
    ponente = row['Ponente']
    ws1.append([hora, titulo, tipo, sala, ponente])
    current_row = ws1.max_row
    for col in range(1, 6):
        cell = ws1.cell(row=current_row, column=col)
        cell.alignment = center_align
        cell.border = border
    if row["Solapado"]:
        ws1.cell(row=current_row, column=1).fill = red_fill
        ws1.cell(row=current_row, column=2).fill = red_fill

ws2 = wb.create_sheet(title="Solapamientos")
ws2.append(["Hora", "Eventos que se solapan"])
for col_num, _ in enumerate(["Hora", "Eventos que se solapan"], 1):
    cell = ws2.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = border

solapados_df = df[df["Solapado"]]
intervals = []
for idx, event in solapados_df.iterrows():
    intervals.append((event["Hora inicio"], event["Hora fin"], event["Título"]))

def agrupar_solapamientos(intervals):
    intervals.sort(key=lambda x: time_to_minutes(x[0]))
    grupos = []
    current_group = [intervals[0]]
    for i in range(1, len(intervals)):
        prev_start, prev_end, _ = current_group[-1]
        curr_start, curr_end, _ = intervals[i]
        if time_to_minutes(curr_start) < time_to_minutes(prev_end):
            current_group.append(intervals[i])
            max_end = max(time_to_minutes(prev_end), time_to_minutes(curr_end))
            current_group[-1] = (curr_start, curr_end, intervals[i][2])
        else:
            grupos.append(current_group)
            current_group = [intervals[i]]
    grupos.append(current_group)
    return grupos

grupos = agrupar_solapamientos(intervals)

for grupo in grupos:
    horas_inicio = [time_to_minutes(x[0]) for x in grupo]
    horas_fin = [time_to_minutes(x[1]) for x in grupo]
    min_inicio = min(horas_inicio)
    max_fin = max(horas_fin)
    hora_texto = f"{min_inicio//60:02d}:{min_inicio%60:02d} - {max_fin//60:02d}:{max_fin%60:02d}"
    eventos_texto = "; ".join([x[2] for x in grupo])
    ws2.append([hora_texto, eventos_texto])
    current_row = ws2.max_row
    for col in range(1, 3):
        cell = ws2.cell(row=current_row, column=col)
        cell.alignment = center_align
        cell.border = border

wb.save("BiznagaFest2025_programa.xlsx")
print("Archivo BiznagaFest2025_programa.xlsx generado con éxito.")
